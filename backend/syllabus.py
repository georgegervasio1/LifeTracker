"""Education: course syllabi — the original file, its text, and the dates that matter.

Data (per user, under users/{uid}/):
  edu_syllabi/{id}        -> structured syllabus: course link, professor & contact, meeting info, grading,
                             policies, class-by-class sessions, deadlines, plus the extracted text
  edu_syllabus_files/{id} -> the original upload (bytes), same id as its syllabus

A syllabus gets its structure either from an import file (POST /api/edu/syllabi/import) or by uploading
the PDF / Word / text file, in which case Claude reads it (if ANTHROPIC_API_KEY is set).
Deadlines and sessions feed the Education calendar via calendar_events().
"""
import io
import json
import os
import re
from datetime import datetime, timezone

from flask import Blueprint, Response, jsonify, request

bp = Blueprint("syllabus", __name__)
db = firestore = require_owner = user_ref = None

MAX_FILE_BYTES = 900_000          # Firestore documents cap at 1 MiB
MAX_TEXT = 120_000
DEADLINE_KINDS = {"assignment", "paper", "presentation", "exam", "reading", "other"}
TYPES = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         "txt": "text/plain", "md": "text/markdown"}


def init(app, _db, _firestore, _require_owner, _user_ref):
    global db, firestore, require_owner, user_ref
    db, firestore, require_owner, user_ref = _db, _firestore, _require_owner, _user_ref
    _register_routes()
    app.register_blueprint(bp)


# ---------------------------------------------------------------- cleaning
def _s(v, n=500):
    return None if v in (None, "") else str(v).strip()[:n]


def _date(v):
    v = _s(v, 10)
    return v if v and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) else None


def clean_syllabus(b):
    """Keep only known fields, in known shapes. Unknown or malformed rows are dropped, not rejected."""
    out = {}
    for k in ("course_title", "term", "code", "professor", "email", "office", "office_hours", "meets", "location",
              "summary", "exam", "textbook", "course_id"):
        if k in b:
            out[k] = _s(b[k], 1500 if k == "summary" else 500)
    if "match" in b:
        out["match"] = [_s(m, 80) for m in (b.get("match") or []) if _s(m, 80)][:12]
    if "grading" in b:
        out["grading"] = [{"item": _s(g.get("item"), 200), "weight": _s(g.get("weight"), 40)}
                          for g in (b.get("grading") or []) if isinstance(g, dict) and g.get("item")][:20]
    if "policies" in b:
        out["policies"] = [_s(p, 600) for p in (b.get("policies") or []) if _s(p)][:20]
    if "sessions" in b:
        rows = []
        for r in (b.get("sessions") or [])[:80]:
            if not isinstance(r, dict) or not (r.get("title") or r.get("reading")):
                continue
            rows.append({"date": _date(r.get("date")), "time": _s(r.get("time"), 40), "title": _s(r.get("title"), 300),
                         "reading": _s(r.get("reading"), 1200), "note": _s(r.get("note"), 600),
                         "no_class": bool(r.get("no_class"))})
        out["sessions"] = sorted(rows, key=lambda r: r["date"] or "9999")
    if "deadlines" in b:
        rows = []
        for r in (b.get("deadlines") or [])[:60]:
            if not isinstance(r, dict) or not r.get("title"):
                continue
            kind = (r.get("kind") or "assignment").lower()
            rows.append({"date": _date(r.get("date")), "time": _s(r.get("time"), 40), "title": _s(r.get("title"), 300),
                         "kind": kind if kind in DEADLINE_KINDS else "other", "note": _s(r.get("note"), 600)})
        out["deadlines"] = sorted(rows, key=lambda r: r["date"] or "9999")
    return out


# ---------------------------------------------------------------- course matching
def _tok(s):
    from education import _tokens
    return _tokens(s)


def match_course(syl, courses):
    """Best course for a syllabus: explicit id, then code, then the most title words in common."""
    if syl.get("course_id") and any(c.get("id") == syl["course_id"] for c in courses):
        return syl["course_id"]
    pool = [c for c in courses if c.get("status") != "dropped" and (not syl.get("term") or c.get("term") == syl["term"])] or courses
    code = (syl.get("code") or "").replace(" ", "").lower()
    if code:
        for c in pool:
            if (c.get("code") or "").replace(" ", "").lower() == code:
                return c.get("id")
    names = [syl.get("course_title") or ""] + list(syl.get("match") or [])
    want = set().union(*[_tok(n) for n in names]) if names else set()
    best, score = None, 0
    for c in pool:
        have = _tok(c.get("title"))
        s = len(want & have) / max(1, len(have))
        if s > score:
            best, score = c.get("id"), s
    return best if score >= 0.5 else None


# ---------------------------------------------------------------- reading the file
def extract_text(name, data):
    ext = (name or "").rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in r.cells) for r in t.rows]
        return "\n".join(parts)
    return data.decode("utf-8", errors="replace")


PARSE_TOOL = {
    "name": "save_syllabus",
    "description": "Save the structured contents of a law school course syllabus.",
    "input_schema": {
        "type": "object",
        "properties": {
            "course_title": {"type": "string"}, "code": {"type": "string"}, "professor": {"type": "string"},
            "email": {"type": "string"}, "office": {"type": "string"}, "office_hours": {"type": "string"},
            "meets": {"type": "string", "description": "Days and times, e.g. 'M 9:00-11:00 am'"},
            "location": {"type": "string"}, "textbook": {"type": "string"},
            "exam": {"type": "string", "description": "Final exam or paper format, in one line"},
            "summary": {"type": "string", "description": "Two sentences on what the course covers and how it is graded"},
            "grading": {"type": "array", "items": {"type": "object", "properties": {
                "item": {"type": "string"}, "weight": {"type": "string"}}, "required": ["item"]}},
            "policies": {"type": "array", "items": {"type": "string"},
                         "description": "Short rules a student must not miss: attendance, laptops, AI, cold-calling, late penalties"},
            "sessions": {"type": "array", "items": {"type": "object", "properties": {
                "date": {"type": "string", "description": "YYYY-MM-DD"}, "time": {"type": "string"},
                "title": {"type": "string"}, "reading": {"type": "string"}, "note": {"type": "string"},
                "no_class": {"type": "boolean"}}}},
            "deadlines": {"type": "array", "items": {"type": "object", "properties": {
                "date": {"type": "string", "description": "YYYY-MM-DD"}, "time": {"type": "string"},
                "title": {"type": "string"}, "kind": {"type": "string", "enum": sorted(DEADLINE_KINDS)},
                "note": {"type": "string"}}, "required": ["title"]}},
        },
    },
}


def parse_with_claude(text, term):
    if not os.environ.get("ANTHROPIC_API_KEY"):
        return {}
    from assistant import MODEL, client
    year = re.search(r"(\d{4})", term or "")
    msg = client().messages.create(
        model=MODEL, max_tokens=8000, tools=[PARSE_TOOL], tool_choice={"type": "tool", "name": "save_syllabus"},
        messages=[{"role": "user", "content":
                   f"This is the syllabus for a {term or 'current'} course. Extract it with save_syllabus. "
                   f"Dates without a year are in {year.group(1) if year else 'the term year'}. One session per class "
                   "meeting with its reading; deadlines are anything due or scheduled outside an ordinary class "
                   "(papers, drafts, presentations, exams, special sessions). Leave out anything the syllabus doesn't say.\n\n"
                   + text[:60_000]}])
    for block in msg.content:
        if getattr(block, "type", "") == "tool_use":
            return block.input
    return {}


# ---------------------------------------------------------------- state + calendar
def load_syllabi(courses):
    out = []
    for d in user_ref().collection("edu_syllabi").stream():
        s = d.to_dict() | {"id": d.id}
        s.pop("text", None)
        if not s.get("course_id") or not any(c.get("id") == s["course_id"] for c in courses):
            s["course_id"] = match_course(s, courses)
        out.append(s)
    out.sort(key=lambda s: (s.get("course_title") or "").lower())
    return out


def calendar_events(syllabi, courses):
    """Syllabus deadlines (and cancelled classes) as Education calendar events, grouped by term."""
    by_id = {c.get("id"): c for c in courses}
    out = {}
    for s in syllabi:
        c = by_id.get(s.get("course_id")) or {}
        term = c.get("term") or s.get("term") or "Other"
        name = s.get("short") or c.get("title") or s.get("course_title") or "Course"
        for d in s.get("deadlines") or []:
            if d.get("date"):
                out.setdefault(term, []).append({"date": d["date"], "end": None, "label": f"{name}: {d['title']}",
                                                 "kind": "assignment", "time": d.get("time"), "course_id": s.get("course_id"),
                                                 "syllabus_id": s["id"]})
        for r in s.get("sessions") or []:
            if r.get("no_class") and r.get("date"):
                out.setdefault(term, []).append({"date": r["date"], "end": None, "label": f"{name}: no class",
                                                 "kind": "other", "course_id": s.get("course_id"), "syllabus_id": s["id"]})
    return out


def _now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


# ---------------------------------------------------------------- routes
def _register_routes():
    from education import education_state, load_all

    def col():
        return user_ref().collection("edu_syllabi")

    @bp.post("/api/edu/syllabi/import")
    @require_owner
    def import_syllabi():
        """Upsert syllabi from JSON {syllabi:[...]}. Matched to an existing syllabus by course title + term;
        an uploaded original file is kept."""
        b = request.get_json(force=True) or {}
        _, courses = load_all()
        existing = {((d.to_dict().get("term") or "").lower(), (d.to_dict().get("course_title") or "").strip().lower()): d.id
                    for d in col().stream()}
        n = 0
        for raw in (b.get("syllabi") or [])[:40]:
            if not isinstance(raw, dict) or not raw.get("course_title"):
                continue
            s = clean_syllabus(raw)
            s["short"] = _s(raw.get("short"), 40)
            cid = match_course(s, courses)
            if cid:
                s["course_id"] = cid
            else:
                s.pop("course_id", None)
            s["updated_at"] = _now()
            key = ((s.get("term") or "").lower(), s["course_title"].strip().lower())
            ref = col().document(existing[key]) if key in existing else col().document()
            ref.set(s, merge=True)
            n += 1
        return jsonify(education_state() | {"imported": n})

    @bp.post("/api/edu/syllabi/upload")
    @require_owner
    def upload_syllabus():
        """Multipart: file, plus syllabus_id (replace that one's file) or course_id (attach to that course).
        Claude reads the file when the syllabus has no class schedule yet, or when reparse=1."""
        f = request.files.get("file")
        if not f or not f.filename:
            return jsonify(error="choose a file"), 400
        ext = f.filename.rsplit(".", 1)[-1].lower()
        if ext not in TYPES:
            return jsonify(error="upload a PDF, Word (.docx) or text file"), 400
        data = f.read()
        if len(data) > MAX_FILE_BYTES:
            return jsonify(error="that file is over 900 KB; save a smaller PDF"), 400
        try:
            text = extract_text(f.filename, data)[:MAX_TEXT]
        except Exception as e:
            return jsonify(error=f"couldn't read that file ({e.__class__.__name__})"), 400

        _, courses = load_all()
        sid, cid = request.form.get("syllabus_id"), request.form.get("course_id")
        course = next((c for c in courses if c.get("id") == cid), None)
        ref = col().document(sid) if sid else None
        cur = ref.get().to_dict() if ref and ref.get().exists else None
        if sid and cur is None:
            return jsonify(error="syllabus not found"), 404
        if not ref and course:      # one syllabus per course: reuse it if there is one
            hit = next((d for d in col().stream() if d.to_dict().get("course_id") == cid), None)
            ref, cur = (hit.reference, hit.to_dict()) if hit else (None, None)
        ref = ref or col().document()

        upd = {"file_name": f.filename[:200], "file_type": TYPES[ext], "file_size": len(data), "has_file": True,
               "text": text, "uploaded_at": _now(), "updated_at": _now()}
        if course:
            upd |= {"course_id": course["id"], "term": course.get("term")}
            if not (cur or {}).get("course_title"):
                upd["course_title"] = course.get("title")
        parsed = False
        if request.form.get("reparse") == "1" or not (cur or {}).get("sessions"):
            try:
                got = clean_syllabus(parse_with_claude(text, upd.get("term") or (cur or {}).get("term")))
                if course:
                    got.pop("course_title", None)
                upd |= got
                parsed = bool(got)
            except Exception as e:      # keep the file even if reading it fails
                upd["parse_error"] = str(e)[:300]
        if not upd.get("course_title") and not (cur or {}).get("course_title"):
            upd["course_title"] = f.filename.rsplit(".", 1)[0][:120]
        ref.set(upd, merge=True)
        user_ref().collection("edu_syllabus_files").document(ref.id).set({"data": data, "content_type": TYPES[ext]})
        return jsonify(education_state() | {"id": ref.id, "parsed": parsed})

    @bp.patch("/api/edu/syllabi/<sid>")
    @require_owner
    def patch_syllabus(sid):
        ref = col().document(sid)
        if not ref.get().exists:
            return jsonify(error="syllabus not found"), 404
        b = request.get_json(force=True) or {}
        upd = clean_syllabus(b)
        upd["updated_at"] = _now()
        ref.set(upd, merge=True)
        return jsonify(education_state())

    @bp.get("/api/edu/syllabi/<sid>/file")
    @require_owner
    def get_file(sid):
        snap = user_ref().collection("edu_syllabus_files").document(sid).get()
        meta = col().document(sid).get()
        if not snap.exists:
            return jsonify(error="no file for this syllabus"), 404
        f = snap.to_dict()
        name = (meta.to_dict() or {}).get("file_name", "syllabus") if meta.exists else "syllabus"
        return Response(bytes(f["data"]), mimetype=f.get("content_type") or "application/octet-stream",
                        headers={"Content-Disposition": f'inline; filename="{name}"'})

    @bp.delete("/api/edu/syllabi/<sid>")
    @require_owner
    def delete_syllabus(sid):
        col().document(sid).delete()
        user_ref().collection("edu_syllabus_files").document(sid).delete()
        return jsonify(education_state())
